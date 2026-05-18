"""
Tests for open_notebook/utils/docx_table_extractor.py

Uses the real test file "VB TEST BANG BIEU.docx" where available,
plus in-memory DOCX documents constructed with python-docx for unit tests.
"""

import os
import pytest

from docx import Document
from docx.oxml.ns import qn

from open_notebook.utils.docx_table_extractor import (
    _table_to_markdown,
    extract_docx_with_tables,
)

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

TESTS_DIR = os.path.dirname(__file__)
REAL_DOCX = os.path.join(TESTS_DIR, "VB TEST BANG BIEU.docx")


def _make_docx_with_table(tmp_path, table_data: list[list[str]], paragraphs_before=None, paragraphs_after=None) -> str:
    """Build a minimal in-memory DOCX with optional surrounding paragraphs."""
    doc = Document()
    if paragraphs_before:
        for text in paragraphs_before:
            doc.add_paragraph(text)
    table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
    for r_idx, row_data in enumerate(table_data):
        for c_idx, cell_text in enumerate(row_data):
            table.rows[r_idx].cells[c_idx].text = cell_text
    if paragraphs_after:
        for text in paragraphs_after:
            doc.add_paragraph(text)
    out_path = str(tmp_path / "test.docx")
    doc.save(out_path)
    return out_path


# ---------------------------------------------------------------------------
# Unit tests: _table_to_markdown
# ---------------------------------------------------------------------------

class TestTableToMarkdown:
    def test_simple_table(self, tmp_path):
        doc = Document()
        table = doc.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "Header 1"
        table.rows[0].cells[1].text = "Header 2"
        table.rows[1].cells[0].text = "Cell A"
        table.rows[1].cells[1].text = "Cell B"
        result = _table_to_markdown(table)
        assert "| --- |" in result
        assert "Header 1" in result
        assert "Cell A" in result

    def test_empty_cells_become_empty_string(self, tmp_path):
        doc = Document()
        table = doc.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "H1"
        table.rows[0].cells[1].text = "H2"
        table.rows[1].cells[0].text = ""
        table.rows[1].cells[1].text = "Value"
        result = _table_to_markdown(table)
        assert "|  |" in result or "| |" in result or "Value" in result

    def test_table_with_no_rows_skipped(self, tmp_path):
        """Tables that yield completely empty headers should return empty string."""
        doc = Document()
        # A table whose first row cells are all blank
        table = doc.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = ""
        table.rows[0].cells[1].text = ""
        result = _table_to_markdown(table)
        assert result == ""

    def test_cell_newlines_replaced_with_space(self, tmp_path):
        doc = Document()
        table = doc.add_table(rows=2, cols=1)
        table.rows[0].cells[0].text = "Header"
        table.rows[1].cells[0].text = "Line1\nLine2"
        result = _table_to_markdown(table)
        assert "\n" not in result.split("---")[1].split("\n")[-1]


# ---------------------------------------------------------------------------
# Integration tests: extract_docx_with_tables
# ---------------------------------------------------------------------------

class TestExtractDocxWithTables:
    def test_simple_table_produces_gfm_separator(self, tmp_path):
        """DOCX with a simple table → output contains | --- |"""
        path = _make_docx_with_table(
            tmp_path,
            [["Name", "Age"], ["Alice", "30"], ["Bob", "25"]],
        )
        result = extract_docx_with_tables(path)
        assert "| --- |" in result

    def test_paragraph_table_paragraph_order(self, tmp_path):
        """Paragraph → Table → Paragraph must appear in that order."""
        path = _make_docx_with_table(
            tmp_path,
            [["Col1", "Col2"], ["A", "B"]],
            paragraphs_before=["Before paragraph"],
            paragraphs_after=["After paragraph"],
        )
        result = extract_docx_with_tables(path)
        before_pos = result.index("Before paragraph")
        table_pos = result.index("| --- |")
        after_pos = result.index("After paragraph")
        assert before_pos < table_pos < after_pos, (
            f"Expected before({before_pos}) < table({table_pos}) < after({after_pos})"
        )

    def test_docx_without_tables_does_not_error(self, tmp_path):
        """DOCX with no tables should return paragraph text without errors."""
        doc = Document()
        doc.add_paragraph("Just a paragraph, no tables.")
        out_path = str(tmp_path / "no_table.docx")
        doc.save(out_path)
        result = extract_docx_with_tables(out_path)
        assert "Just a paragraph" in result
        assert "| --- |" not in result

    @pytest.mark.skipif(
        not os.path.exists(REAL_DOCX),
        reason="Real test DOCX not found at tests/VB TEST BANG BIEU.docx",
    )
    def test_real_test_file_contains_tables(self):
        """The actual VB TEST BANG BIEU.docx should produce at least one table."""
        result = extract_docx_with_tables(REAL_DOCX)
        assert "| --- |" in result, (
            "Expected at least one Markdown table separator in real test file output"
        )
